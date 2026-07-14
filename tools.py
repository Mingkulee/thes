"""External tools the agents can call."""

from __future__ import annotations

import io
import logging
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from typing import Any
from xml.etree import ElementTree as ET

import httpx

logger = logging.getLogger(__name__)

_GEOCODE_URL = "https://geocoding-api.open-meteo.com/v1/search"
_FORECAST_URL = "https://api.open-meteo.com/v1/forecast"


async def geocode_location(
    location: str,
) -> tuple[float, float, str, str] | None:
    """Resolve a place name to (latitude, longitude, display_name, country).

    Returns None if no match was found.
    """
    async with httpx.AsyncClient(timeout=10.0) as client:
        resp = await client.get(
            _GEOCODE_URL,
            params={"name": location, "count": 1, "language": "ko"},
        )
        resp.raise_for_status()
        data = resp.json()

    results = data.get("results") or []
    if not results:
        return None

    r = results[0]
    return (
        float(r["latitude"]),
        float(r["longitude"]),
        r.get("name", location),
        r.get("country", ""),
    )


async def fetch_weather(latitude: float, longitude: float) -> dict[str, Any]:
    """Fetch current weather and a 3-day forecast from Open-Meteo."""
    params = {
        "latitude": latitude,
        "longitude": longitude,
        "current": ",".join(
            [
                "temperature_2m",
                "apparent_temperature",
                "relative_humidity_2m",
                "weather_code",
                "wind_speed_10m",
                "precipitation",
            ]
        ),
        "daily": ",".join(
            [
                "weather_code",
                "temperature_2m_max",
                "temperature_2m_min",
                "precipitation_probability_max",
                "precipitation_sum",
            ]
        ),
        "timezone": "auto",
        "forecast_days": 3,
    }

    async with httpx.AsyncClient(timeout=10.0) as client:
        resp = await client.get(_FORECAST_URL, params=params)
        resp.raise_for_status()
        return resp.json()


# Open-Meteo WMO weather codes → 한국어 설명 (간략 매핑)
WEATHER_CODE_KO: dict[int, str] = {
    0: "맑음",
    1: "대체로 맑음",
    2: "부분적으로 흐림",
    3: "흐림",
    45: "안개",
    48: "서리 안개",
    51: "약한 이슬비",
    53: "보통 이슬비",
    55: "강한 이슬비",
    56: "약한 어는 이슬비",
    57: "강한 어는 이슬비",
    61: "약한 비",
    63: "보통 비",
    65: "강한 비",
    66: "약한 어는 비",
    67: "강한 어는 비",
    71: "약한 눈",
    73: "보통 눈",
    75: "강한 눈",
    77: "싸락눈",
    80: "약한 소나기",
    81: "보통 소나기",
    82: "강한 소나기",
    85: "약한 눈 소나기",
    86: "강한 눈 소나기",
    95: "뇌우",
    96: "약한 우박을 동반한 뇌우",
    99: "강한 우박을 동반한 뇌우",
}


def describe_weather_code(code: int) -> str:
    return WEATHER_CODE_KO.get(code, f"날씨 코드 {code}")


# ---------------------------------------------------------------------------
# HWP / HWPX document extraction
# ---------------------------------------------------------------------------

_HP = "http://www.hancom.co.kr/hwpml/2012/paragraph"


def extract_document(file_bytes: bytes, filename: str) -> dict[str, Any]:
    """Extract text, tables, and image count from an HWP or HWPX file."""
    fl = filename.lower()
    if fl.endswith(".hwpx"):
        return _extract_hwpx(file_bytes)
    if fl.endswith(".hwp"):
        return _extract_hwp(file_bytes)
    raise ValueError(f"지원하지 않는 형식: {filename}")


# ── HWPX (ZIP + XML) ────────────────────────────────────────────────────────

def _extract_hwpx(file_bytes: bytes) -> dict[str, Any]:
    paragraphs: list[str] = []
    tables: list[list[list[str]]] = []
    image_count = 0

    with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
        names = zf.namelist()
        image_count = sum(
            1 for n in names
            if re.search(r"\.(png|jpe?g|gif|bmp|tiff?|emf|wmf)$", n, re.I)
        )
        section_files = sorted(
            n for n in names if re.match(r"Contents/section\d+\.xml", n)
        )
        if not section_files:
            section_files = sorted(
                n for n in names
                if "section" in n.lower() and n.endswith(".xml")
            )
        for sf in section_files:
            _parse_hwpx_section(zf.read(sf), paragraphs, tables)

    return {
        "format": "HWPX",
        "text": "\n".join(p for p in paragraphs if p.strip()),
        "tables": tables,
        "image_count": image_count,
    }


def _parse_hwpx_section(
    xml_bytes: bytes,
    paragraphs: list[str],
    tables: list[list[list[str]]],
) -> None:
    try:
        root = ET.fromstring(xml_bytes)
    except ET.ParseError:
        return

    T = f"{{{_HP}}}t"
    P = f"{{{_HP}}}p"
    TBL = f"{{{_HP}}}tbl"
    TR = f"{{{_HP}}}tr"
    TC = f"{{{_HP}}}tc"

    def para_text(p_elem: ET.Element) -> str:
        return "".join(t.text or "" for t in p_elem.iter(T))

    def parse_table(tbl_elem: ET.Element) -> list[list[str]]:
        rows = []
        for tr in tbl_elem.iter(TR):
            row = []
            for tc in tr.findall(TC):
                cell = " ".join(
                    para_text(p).strip()
                    for p in tc.iter(P)
                    if para_text(p).strip()
                )
                row.append(cell)
            if row:
                rows.append(row)
        return rows

    def walk(elem: ET.Element) -> None:
        for child in elem:
            if child.tag == TBL:
                tbl = parse_table(child)
                if tbl:
                    tables.append(tbl)
                    paragraphs.append(f"[표 {len(tables)}]")
            elif child.tag == P:
                paragraphs.append(para_text(child))
            else:
                walk(child)

    walk(root)


# ── HWP binary ──────────────────────────────────────────────────────────────

def _extract_hwp(file_bytes: bytes) -> dict[str, Any]:
    result = _try_hwp5txt(file_bytes)
    if result is not None:
        return result
    result = _try_libreoffice(file_bytes, ".hwp")
    if result is not None:
        return result
    return {
        "format": "HWP",
        "text": (
            "(HWP 바이너리 파싱 실패)\n"
            "pyhwp 또는 LibreOffice가 설치되어 있지 않습니다.\n"
            "pip install pyhwp  또는  sudo apt-get install libreoffice"
        ),
        "tables": [],
        "image_count": 0,
    }


def _try_hwp5txt(file_bytes: bytes) -> dict[str, Any] | None:
    if not shutil.which("hwp5txt"):
        return None
    tmp = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".hwp", delete=False) as f:
            f.write(file_bytes)
            tmp = f.name
        proc = subprocess.run(
            ["hwp5txt", tmp],
            capture_output=True,
            text=True,
            timeout=30,
        )
        if proc.returncode == 0 and proc.stdout.strip():
            return {
                "format": "HWP (pyhwp)",
                "text": proc.stdout.strip(),
                "tables": [],
                "image_count": 0,
            }
    except Exception:
        logger.exception("hwp5txt failed")
    finally:
        if tmp and os.path.exists(tmp):
            os.unlink(tmp)
    return None


def _try_libreoffice(file_bytes: bytes, suffix: str) -> dict[str, Any] | None:
    lo = shutil.which("libreoffice") or shutil.which("soffice")
    if not lo:
        return None
    try:
        with tempfile.TemporaryDirectory() as tmpdir:
            src = os.path.join(tmpdir, f"input{suffix}")
            with open(src, "wb") as f:
                f.write(file_bytes)
            subprocess.run(
                [lo, "--headless", "--convert-to", "docx", "--outdir", tmpdir, src],
                capture_output=True,
                timeout=60,
            )
            docx_path = os.path.join(tmpdir, "input.docx")
            if os.path.exists(docx_path):
                with open(docx_path, "rb") as f:
                    return _extract_docx(f.read())
    except Exception:
        logger.exception("libreoffice conversion failed")
    return None


def _extract_docx(docx_bytes: bytes) -> dict[str, Any]:
    from docx import Document  # python-docx (optional dep)

    doc = Document(io.BytesIO(docx_bytes))
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    tables = []
    for tbl in doc.tables:
        tables.append([[cell.text.strip() for cell in row.cells] for row in tbl.rows])
    image_count = sum(
        1 for rel in doc.part.rels.values() if "image" in rel.reltype
    )
    return {
        "format": "HWP→DOCX (LibreOffice)",
        "text": "\n".join(paras),
        "tables": tables,
        "image_count": image_count,
    }


# ---------------------------------------------------------------------------
# Gmail API tools
# ---------------------------------------------------------------------------

import asyncio as _asyncio
import base64


def _extract_email_body(payload: dict) -> str:
    """Gmail API payload에서 텍스트 본문을 재귀적으로 추출합니다."""
    mime = payload.get("mimeType", "")
    body_data = payload.get("body", {}).get("data", "")

    if body_data:
        text = base64.urlsafe_b64decode(body_data).decode("utf-8", errors="replace")
        if mime == "text/html":
            text = re.sub(r"<[^>]+>", " ", text)
            text = re.sub(r"\s{2,}", " ", text).strip()
        return text

    for part in payload.get("parts", []):
        if part.get("mimeType") == "text/plain":
            data = part.get("body", {}).get("data", "")
            if data:
                return base64.urlsafe_b64decode(data).decode("utf-8", errors="replace")

    for part in payload.get("parts", []):
        result = _extract_email_body(part)
        if result:
            return result

    return payload.get("snippet", "")


def _fetch_gmail_sync(max_results: int, query: str) -> list[dict[str, Any]]:
    from gmail_auth import get_gmail_credentials
    from googleapiclient.discovery import build  # type: ignore

    creds = get_gmail_credentials()
    service = build("gmail", "v1", credentials=creds)

    result = service.users().messages().list(
        userId="me", q=query, maxResults=max_results
    ).execute()

    messages = []
    for ref in result.get("messages", []):
        msg = service.users().messages().get(
            userId="me", id=ref["id"], format="full"
        ).execute()
        headers = {h["name"]: h["value"] for h in msg["payload"].get("headers", [])}
        body = _extract_email_body(msg["payload"])
        messages.append({
            "id": msg["id"],
            "from": headers.get("From", ""),
            "subject": headers.get("Subject", "(제목 없음)"),
            "date": headers.get("Date", ""),
            "snippet": msg.get("snippet", ""),
            "body": body[:3000],
        })
    return messages


async def fetch_gmail_messages(
    max_results: int = 5, query: str = "is:unread"
) -> list[dict[str, Any]]:
    """Gmail에서 메일 목록을 가져옵니다 (비동기 래퍼)."""
    return await _asyncio.to_thread(_fetch_gmail_sync, max_results, query)
